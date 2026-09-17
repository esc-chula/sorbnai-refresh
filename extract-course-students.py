"""
Extract enrolled student IDs from per-course roster files and combine them with
sheet-range.json room/range data.

Usage:
    py extract-course-students.py
    py extract-course-students.py --roster-dir src/data --schedule sheet-range.json --output course-students.json

Supported roster sources:
    - .xlsx: extracts 10-digit student IDs from workbook cells. When the
      workbook contains exam-room blocks, those blocks are used as the source
      of truth for room and seat order.

The filename should contain the course code, for example:
    5500111.xlsx
"""

import argparse
import glob
import json
import os
import re
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl


DEFAULT_ROSTER_DIR = os.path.join("src", "data")
DEFAULT_SCHEDULE = "sheet-range.json"
DEFAULT_OUTPUT = "course-students.json"
STUDENT_ID_RE = re.compile(r"\b\d{10}\b")
SPACED_STUDENT_ID_RE = re.compile(r"\b(\d{3})\s+(\d{5})\s+(\d{2})\b")


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def course_code_from_filename(path: str) -> Optional[str]:
    match = re.search(r"(\d{7})", os.path.basename(path))
    return match.group(1) if match else None


def extract_ids_from_text(text: str) -> List[str]:
    ids = set(STUDENT_ID_RE.findall(text or ""))
    for match in SPACED_STUDENT_ID_RE.finditer(text or ""):
        ids.add("".join(match.groups()))
    return sorted(ids)


def extract_xlsx_student_records(path: str) -> List[Dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    records: List[Dict[str, Any]] = []
    seen_ids: Set[str] = set()

    for sheet in workbook.worksheets:
        for cell_row in sheet.iter_rows():
            row = tuple(cell.value for cell in cell_row)
            if not row or not is_student_row(row):
                continue
            student_id = student_id_from_row(row[1:4])
            if not student_id or student_id in seen_ids:
                continue
            seen_ids.add(student_id)
            records.append(
                {
                    "id": student_id,
                    "name": clean(row[2]) if len(row) > 2 else "",
                    "seat": seat_from_row(row),
                }
            )

    return records


def room_key(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z]+", "", clean(value)).upper()


def building_key(value: Any) -> str:
    return re.sub(r"[^0-9A-Za-z]+", "", clean(value)).upper()


def normalize_header_building(value: str) -> str:
    text = clean(value).upper().replace(" ", "")
    if text == "100":
        return "EN 100"
    if re.fullmatch(r"\d+", text):
        return "ENG {}".format(text)
    if text.startswith("ENG"):
        return "ENG {}".format(text[3:]) if text[3:].isdigit() else text
    return clean(value)


def parse_block_room_header(value: Any) -> Optional[Dict[str, str]]:
    text = clean(value)
    if not text or "/" not in text:
        return None

    match = re.fullmatch(r"([0-9A-Za-z]+(?:-\d+)?)\s*/\s*([0-9A-Za-z()ก-๙ ]+)", text)
    if not match:
        return None

    building_part, room_part = match.groups()
    if "-" in building_part:
        building_raw, room_prefix = building_part.split("-", 1)
        room = "{}/{}".format(room_prefix, clean(room_part))
    else:
        building_raw = building_part
        room = clean(room_part)

    return {
        "building": normalize_header_building(building_raw),
        "room": room,
    }


def parse_thai_room_header(value: Any) -> Optional[Dict[str, str]]:
    text = clean(value)
    if not text or "ห้อง" not in text:
        return None

    room_match = re.search(r"ห้อง\s*([0-9A-Za-z/()ก-๙ -]+?)(?:\s+ตึก|\s+อาคาร|$)", text)
    building_match = re.search(r"(?:ตึก|อาคาร)\s*([0-9A-Za-z ]+)", text)
    if not room_match:
        return None

    building = clean(building_match.group(1)) if building_match else ""
    if building_key(building).startswith("ENG"):
        suffix = building_key(building)[3:]
        building = "ENG {}".format(suffix) if suffix else "ENG"

    return {
        "building": building,
        "room": clean(room_match.group(1)),
    }


def schedule_group_for_header(
    header: Dict[str, str],
    schedule_groups: List[Dict[str, Any]],
    used_group_indexes: Set[int],
) -> Dict[str, Any]:
    header_room = room_key(header.get("room"))
    header_building = building_key(header.get("building"))

    for index, group in enumerate(schedule_groups):
        if index in used_group_indexes:
            continue
        if room_key(group.get("room")) != header_room:
            continue
        if header_building and building_key(group.get("building")) != header_building:
            continue
        used_group_indexes.add(index)
        return group

    for index, group in enumerate(schedule_groups):
        if index in used_group_indexes:
            continue
        if room_key(group.get("room")) == header_room:
            used_group_indexes.add(index)
            return group

    return {
        "building": header.get("building", ""),
        "room": header.get("room", ""),
        "students": 0,
        "range": "",
    }


def student_id_from_row(row: Tuple[Any, ...]) -> str:
    for value in row:
        for student_id in extract_ids_from_text(clean(value)):
            return student_id
    return ""


def seat_from_row(row: Tuple[Any, ...]) -> Optional[int]:
    if not row:
        return None
    seat = clean(row[0])
    return int(seat) if re.fullmatch(r"\d+", seat) else None


def is_student_row(row: Tuple[Any, ...]) -> bool:
    if seat_from_row(row) is None:
        return False
    return bool(student_id_from_row(row[1:4]))


def extract_room_groups_from_xlsx(path: str, schedule_course: Dict[str, Any]) -> List[Dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    schedule_groups = schedule_course.get("group", [])
    used_group_indexes: Set[int] = set()
    groups: List[Dict[str, Any]] = []
    current_group: Optional[Dict[str, Any]] = None

    for sheet in workbook.worksheets:
        for cell_row in sheet.iter_rows():
            row = tuple(cell.value for cell in cell_row)
            first_cell = clean(row[0]) if row else ""
            header = parse_block_room_header(first_cell) or parse_thai_room_header(first_cell)

            if header:
                schedule_group = schedule_group_for_header(header, schedule_groups, used_group_indexes)
                current_group = {
                    "building": clean(schedule_group.get("building")) or header.get("building", ""),
                    "room": clean(schedule_group.get("room")) or header.get("room", ""),
                    "students": int(schedule_group.get("students", 0) or 0),
                    "range": clean(schedule_group.get("range")),
                    "student_ids": [],
                    "student_seats": [],
                }
                groups.append(current_group)
                continue

            if not current_group or not is_student_row(row):
                continue

            student_id = student_id_from_row(row[1:4])
            if student_id and student_id not in current_group["student_ids"]:
                current_group["student_ids"].append(student_id)
                current_group["student_seats"].append(
                    {
                        "id": student_id,
                        "name": clean(row[2]) if len(row) > 2 else "",
                        "seat": seat_from_row(row),
                    }
                )

    return [
        group
        for group in groups
        if group["student_ids"]
    ]


def extract_roster_file(path: str, schedule_course: Dict[str, Any]) -> Tuple[str, List[str], List[Dict[str, Any]], List[Dict[str, Any]], Optional[str]]:
    course_code = course_code_from_filename(path)
    if not course_code:
        return "", [], [], [], "filename does not contain a 7-digit course code"

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".xlsx":
            groups = extract_room_groups_from_xlsx(path, schedule_course)
            ids = sorted({
                student_id
                for group in groups
                for student_id in group.get("student_ids", [])
            })
            if groups:
                records = [
                    record
                    for group in groups
                    for record in group.get("student_seats", [])
                ]
                return course_code, ids, groups, records, None
            records = extract_xlsx_student_records(path)
            return course_code, [record["id"] for record in records], [], records, None
        return course_code, [], [], [], "unsupported file type"
    except Exception as exc:
        return course_code, [], [], [], str(exc)


def load_schedule_by_code(path: str) -> Dict[str, Dict[str, Any]]:
    with open(path, "r", encoding="utf-8") as file:
        schedule_by_no = json.load(file)

    by_code: Dict[str, Dict[str, Any]] = {}
    for schedule_no, course in schedule_by_no.items():
        code = clean(course.get("code"))
        if not code:
            continue
        if code not in by_code:
            by_code[code] = {
                "schedule_no": schedule_no,
                "code": code,
                "title": clean(course.get("title")),
                "date": clean(course.get("date")),
                "time": clean(course.get("time")),
                "sum_student": course.get("sum_student", 0),
                "group": [],
            }
        by_code[code]["group"].extend(course.get("group", []))

    return by_code


def parse_range_parts(range_text: str) -> List[Tuple[int, int]]:
    parts: List[Tuple[int, int]] = []
    text = clean(range_text)
    if not text:
        return parts

    for segment in re.split(r"[,;\n]+", text):
        segment = clean(segment)
        if not segment:
            continue

        range_match = re.fullmatch(r"(\d{10})\s*-\s*(\d{10})", segment)
        if range_match:
            start = int(range_match.group(1))
            end = int(range_match.group(2))
            if start > end:
                start, end = end, start
            parts.append((start, end))
            continue

        for exact in STUDENT_ID_RE.findall(segment):
            value = int(exact)
            parts.append((value, value))

    return parts


def id_in_range(student_id: str, range_parts: List[Tuple[int, int]]) -> bool:
    value = int(student_id)
    return any(start <= value <= end for start, end in range_parts)


def assign_students_to_groups(
    course: Dict[str, Any],
    student_ids: List[str],
    student_records: Dict[str, Dict[str, Any]],
) -> Tuple[List[Dict[str, Any]], List[str]]:
    # The exam range sheet is already sorted by student ID and gives the exact
    # room counts. Sequential slicing avoids over-matching ranges such as
    # "6732039721 - 6832040221" that span multiple ID prefixes.
    cursor = 0
    groups = []
    schedule_groups = course.get("group", [])

    for group_index, group in enumerate(schedule_groups):
        range_parts = parse_range_parts(clean(group.get("range")))
        expected_count = int(group.get("students", 0) or 0)

        if expected_count > 0:
            group_student_ids = student_ids[cursor:cursor + expected_count]
            cursor += len(group_student_ids)
        else:
            group_student_ids = [
                student_id
                for student_id in student_ids[cursor:]
                if id_in_range(student_id, range_parts)
            ]
            cursor += len(group_student_ids)

        if group_index == len(schedule_groups) - 1 and cursor < len(student_ids):
            remaining_capacity = max(expected_count - len(group_student_ids), 0)
            if remaining_capacity:
                extra = student_ids[cursor:cursor + remaining_capacity]
                group_student_ids.extend(extra)
                cursor += len(extra)

        groups.append(
            {
                "building": clean(group.get("building")),
                "room": clean(group.get("room")),
                "students": group.get("students", 0),
                "range": clean(group.get("range")),
                "student_ids": group_student_ids,
                "student_seats": [
                    student_records[student_id]
                    for student_id in group_student_ids
                    if student_id in student_records
                ],
            }
        )

    unmatched = student_ids[cursor:]
    return groups, unmatched


def collect_roster_files(roster_dir: str) -> List[str]:
    patterns = [
        os.path.join(roster_dir, "*.xlsx"),
    ]
    files: List[str] = []
    for pattern in patterns:
        files.extend(glob.glob(pattern))

    # Exclude the master exam-room workbook. It is schedule/range data, not a per-course roster.
    return sorted(
        path
        for path in files
        if course_code_from_filename(path) is not None
    )


def build_output(roster_dir: str, schedule_path: str) -> Dict[str, Any]:
    schedule_by_code = load_schedule_by_code(schedule_path)
    roster_files = collect_roster_files(roster_dir)
    extracted_by_code: Dict[str, Dict[str, Any]] = {}
    warnings = []

    for path in roster_files:
        code = course_code_from_filename(path) or ""
        schedule_course = schedule_by_code.get(code, {
            "code": code,
            "title": "",
            "date": "",
            "time": "",
            "sum_student": 0,
            "group": [],
        })
        code, ids, groups, records, error = extract_roster_file(path, schedule_course)
        if not code:
            warnings.append({"file": path, "warning": error})
            continue

        if code not in extracted_by_code:
            extracted_by_code[code] = {
                "source_files": [],
                "student_ids": set(),
                "groups": [],
                "student_records": {},
                "errors": [],
            }

        extracted_by_code[code]["source_files"].append(path)
        extracted_by_code[code]["student_ids"].update(ids)
        extracted_by_code[code]["groups"].extend(groups)
        extracted_by_code[code]["student_records"].update(
            {record["id"]: record for record in records if record.get("id")}
        )
        for group in groups:
            extracted_by_code[code]["student_records"].update(
                {record["id"]: record for record in group.get("student_seats", []) if record.get("id")}
            )
        if error:
            extracted_by_code[code]["errors"].append({"file": path, "error": error})

    courses: Dict[str, Any] = {}
    for code, extracted in sorted(extracted_by_code.items()):
        student_records = extracted["student_records"]
        student_ids = list(student_records)
        student_ids.extend(
            sorted(set(extracted["student_ids"]) - set(student_ids))
        )
        schedule_course = schedule_by_code.get(code, {
            "code": code,
            "title": "",
            "date": "",
            "time": "",
            "sum_student": 0,
            "group": [],
        })
        if extracted.get("groups"):
            groups = extracted["groups"]
            assigned = {
                student_id
                for group in groups
                for student_id in group.get("student_ids", [])
            }
            unmatched = [student_id for student_id in student_ids if student_id not in assigned]
        else:
            groups, unmatched = assign_students_to_groups(
                schedule_course,
                student_ids,
                student_records,
            )

        courses[code] = {
            "code": code,
            "title": schedule_course.get("title", ""),
            "date": schedule_course.get("date", ""),
            "time": schedule_course.get("time", ""),
            "sum_student": schedule_course.get("sum_student", 0),
            "source_files": extracted["source_files"],
            "student_ids": student_ids,
            "student_count": len(student_ids),
            "groups": groups,
            "unmatched_student_ids": unmatched,
            "errors": extracted["errors"],
        }

    return {
        "schedule_source": schedule_path,
        "roster_dir": roster_dir,
        "course_count": len(courses),
        "courses": courses,
        "warnings": warnings,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--roster-dir", default=DEFAULT_ROSTER_DIR)
    parser.add_argument("--schedule", default=DEFAULT_SCHEDULE)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    result = build_output(args.roster_dir, args.schedule)
    with open(args.output, "w", encoding="utf-8") as file:
        json.dump(result, file, ensure_ascii=False, indent=2)
        file.write("\n")

    total_students = sum(
        course["student_count"] for course in result["courses"].values()
    )
    total_errors = sum(
        len(course["errors"]) for course in result["courses"].values()
    )
    print(
        f"Done! {args.output} written with "
        f"{result['course_count']} course(s), {total_students} student id(s), "
        f"{total_errors} file error(s)."
    )


if __name__ == "__main__":
    main()
