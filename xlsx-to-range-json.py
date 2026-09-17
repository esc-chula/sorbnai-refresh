"""
Excel to JSON converter that mirrors csv-to-json.py output.

Usage:
    py xlsx-to-range-json.py
    py xlsx-to-range-json.py input.xlsx
    py xlsx-to-range-json.py input.xlsx output.json

Default input:
    src/data/(มีรหัสนิสิต) - ตารางรายวิชาที่จะจัดสอบ Midterm ภาคต้น 69.xlsx

Output shape:
{
  "1": {
    "code": "5500111",
    "title": "EXP ENG I",
    "date": "MON 21 SEP 26",
    "time": "08:30-10:30",
    "sum_student": 814,
    "group": [
      {
        "building": "ENG 3",
        "room": "315",
        "students": 103,
        "range": "6830155521 - 6831116221, 6930001621 - 6930106921"
      }
    ]
  }
}
"""

import json
import os
import re
import sys
from typing import Any, Dict, Optional, Tuple

import openpyxl


DEFAULT_INPUT = os.path.join(
    "src",
    "data",
    "(มีรหัสนิสิต) - ตารางรายวิชาที่จะจัดสอบ Midterm ภาคต้น 69.xlsx",
)
DEFAULT_OUTPUT = "sheet.json"

HEADER_ALIASES = {
    "no": ["ที่"],
    "code": ["รหัสวิชา"],
    "title": ["ชื่อย่อรายวิชา"],
    "date": ["วัน", "วันสอบ"],
    "time": ["เวลา", "เวลาสอบ"],
    "sum_student": ["จำนวนนิสิต"],
    "building": ["อาคาร"],
    "room": ["ห้องเรียน"],
    "students": ["จำนวน"],
    "range": ["รหัสนิสิต", "เลขประจำตัวนิสิต"],
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def to_int(value: Any) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    nums = re.findall(r"\d+", str(value))
    return int(nums[-1]) if nums else None


def fmt_time(value: Any) -> str:
    text = clean(value).replace("–", "-").replace(" ", "")
    match = re.fullmatch(r"(\d{4})-(\d{4})", text)
    if match:
        start, end = match.groups()
        return f"{start[:2]}:{start[2:]}-{end[:2]}:{end[2:]}"
    match = re.fullmatch(r"\d{2}:\d{2}-\d{2}:\d{2}", text)
    if match:
        return text
    return clean(value)


def find_header_row(sheet) -> Tuple[int, Dict[str, int]]:
    for row_index in range(1, min(sheet.max_row, 30) + 1):
        header_to_column = {
            clean(sheet.cell(row_index, column).value): column
            for column in range(1, sheet.max_column + 1)
        }
        columns = {}
        for key, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                if alias in header_to_column:
                    columns[key] = header_to_column[alias]
                    break

        if "code" in columns and "range" in columns:
            return row_index, columns

    raise ValueError("Could not find the expected Thai header row.")


def read_row(sheet, row_index: int, columns: Dict[str, int]) -> Dict[str, Any]:
    return {
        key: sheet.cell(row_index, column).value
        for key, column in columns.items()
    }


def is_new_course(row: Dict[str, Any]) -> bool:
    no_value = clean(row.get("no"))
    code = clean(row.get("code"))
    return no_value.isdigit() and code.isdigit()


def append_group(course: Dict[str, Any], row: Dict[str, Any], last_building: str) -> str:
    building = clean(row.get("building")) or last_building
    room = clean(row.get("room"))
    students = to_int(row.get("students"))
    student_range = clean(row.get("range"))

    if not building and not room and not student_range and students is None:
        return last_building

    course["group"].append(
        {
            "building": building,
            "room": room,
            "students": students if students is not None else 0,
            "range": student_range,
        }
    )

    return building or last_building


def convert_xlsx(input_path: str) -> Dict[str, Any]:
    workbook = openpyxl.load_workbook(input_path, data_only=True)
    sheet = workbook.active
    header_row, columns = find_header_row(sheet)

    result = {}
    current_no = ""
    last_building = ""

    for row_index in range(header_row + 1, sheet.max_row + 1):
        row = read_row(sheet, row_index, columns)

        if not any(clean(value) for value in row.values()):
            continue

        if is_new_course(row):
            current_no = clean(row.get("no"))
            last_building = ""

            if current_no not in result:
                sum_students = to_int(row.get("sum_student"))
                result[current_no] = {
                    "code": clean(row.get("code")),
                    "title": clean(row.get("title")),
                    "date": clean(row.get("date")),
                    "time": fmt_time(row.get("time")),
                    "sum_student": sum_students if sum_students is not None else 0,
                    "group": [],
                }

            last_building = append_group(result[current_no], row, last_building)
            continue

        if not current_no:
            continue

        # Skip subtotal rows that only repeat the total count.
        if (
            not clean(row.get("building"))
            and not clean(row.get("room"))
            and not clean(row.get("range"))
        ):
            continue

        last_building = append_group(result[current_no], row, last_building)

    return result


def main() -> None:
    input_path = sys.argv[1] if len(sys.argv) >= 2 else DEFAULT_INPUT
    output_path = sys.argv[2] if len(sys.argv) >= 3 else DEFAULT_OUTPUT

    if len(sys.argv) > 3:
        print("Usage: py xlsx-to-range-json.py [input.xlsx] [output.json]")
        sys.exit(1)

    data = convert_xlsx(input_path)
    with open(output_path, "w", encoding="utf-8") as output:
        json.dump(data, output, ensure_ascii=False, indent=2)
        output.write("\n")

    total_courses = len(data)
    total_groups = sum(len(course["group"]) for course in data.values())
    print(
        f"Done! {output_path} written with "
        f"{total_courses} course(s), {total_groups} group row(s)."
    )


if __name__ == "__main__":
    main()
